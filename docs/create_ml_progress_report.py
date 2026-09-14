"""Build a professor-facing Smart Sleep Alarm ML progress report."""

from pathlib import Path

import matplotlib.pyplot as plt
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUTPUT = Path("docs/Smart-Sleep-Alarm-ML-Progress-Report-2026-09-04.docx")
CHART = Path("ml/artifacts/professor_report_validation_f1.png")
BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
LIGHT_BLUE = "E8EEF5"
LIGHT_GRAY = "F2F4F7"


def set_cell_shading(cell, fill):
    properties = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    properties.append(shading)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    properties = cell._tc.get_or_add_tcPr()
    margins = properties.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        properties.append(margins)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = margins.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_widths(table, widths):
    table.autofit = False
    table_properties = table._tbl.tblPr
    table_width = table_properties.first_child_found_in("w:tblW")
    table_width.set(qn("w:w"), "9360")
    table_width.set(qn("w:type"), "dxa")
    indent = OxmlElement("w:tblInd")
    indent.set(qn("w:w"), "120")
    indent.set(qn("w:type"), "dxa")
    table_properties.append(indent)
    for row in table.rows:
        for cell, width in zip(row.cells, widths):
            cell.width = Inches(width)
            cell_properties = cell._tc.get_or_add_tcPr()
            tc_width = cell_properties.first_child_found_in("w:tcW")
            tc_width.set(qn("w:w"), str(round(width * 1440)))
            tc_width.set(qn("w:type"), "dxa")
            set_cell_margins(cell)


def add_bullet(doc, text):
    paragraph = doc.add_paragraph(style="List Bullet")
    paragraph.add_run(text)
    return paragraph


def add_heading(doc, text, level=1):
    return doc.add_heading(text, level=level)


def add_table(doc, headers, rows, widths):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    set_table_widths(table, widths)
    for cell, text in zip(table.rows[0].cells, headers):
        set_cell_shading(cell, LIGHT_BLUE)
        paragraph = cell.paragraphs[0]
        run = paragraph.add_run(text)
        run.bold = True
        run.font.color.rgb = RGBColor.from_string(DARK_BLUE)
    for number, row_values in enumerate(rows):
        cells = table.add_row().cells
        for cell, value in zip(cells, row_values):
            if number % 2:
                set_cell_shading(cell, "FAFBFC")
            cell.paragraphs[0].add_run(str(value))
    doc.add_paragraph()
    return table


def create_chart():
    CHART.parent.mkdir(parents=True, exist_ok=True)
    names = ["Raw CNN-GRU\n(diagnostic)", "Raw logistic", "Rolling\nengineered", "Causal-time\nengineered", "Hybrid\nvalidation candidate"]
    values = [0.6043, 0.6175, 0.6314, 0.6331, 0.6350]
    colors = ["#AAB7C4", "#7FA8C9", "#5A94C1", "#2E74B5", "#1F4D78"]
    figure, axis = plt.subplots(figsize=(8.2, 4.1))
    bars = axis.bar(names, values, color=colors, width=0.66)
    axis.set_ylim(0.55, 0.65)
    axis.set_ylabel("Best validation F1")
    axis.set_title("Validation comparison (participant-held-out)")
    axis.grid(axis="y", alpha=0.25)
    axis.set_axisbelow(True)
    for bar, value in zip(bars, values):
        axis.text(bar.get_x() + bar.get_width() / 2, value + 0.0015, f"{value:.4f}", ha="center", va="bottom", fontsize=9)
    figure.tight_layout()
    figure.savefig(CHART, dpi=180)
    plt.close(figure)


def configure_styles(doc):
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10
    for style_name, size, color, before, after in (("Heading 1", 16, BLUE, 16, 8), ("Heading 2", 13, BLUE, 12, 6), ("Heading 3", 12, DARK_BLUE, 8, 4)):
        style = styles[style_name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
    styles["List Bullet"].paragraph_format.space_after = Pt(4)
    styles["List Bullet"].paragraph_format.line_spacing = 1.167


def build_report():
    create_chart()
    doc = Document()
    section = doc.sections[0]
    section.top_margin = section.bottom_margin = section.left_margin = section.right_margin = Inches(1)
    section.header_distance = section.footer_distance = Inches(0.492)
    configure_styles(doc)
    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = header.add_run("Smart Sleep Alarm | ML progress brief")
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor(100, 100, 100)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run("Prepared for professor discussion | 4 September 2026").font.size = Pt(9)

    title = doc.add_paragraph()
    title.paragraph_format.space_after = Pt(4)
    title_run = title.add_run("SMART SLEEP ALARM")
    title_run.bold = True
    title_run.font.name = "Calibri"
    title_run.font.size = Pt(23)
    title_run.font.color.rgb = RGBColor.from_string(DARK_BLUE)
    subtitle = doc.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(16)
    subtitle_run = subtitle.add_run("Machine-learning progress, findings, and guidance request")
    subtitle_run.font.size = Pt(14)
    subtitle_run.font.color.rgb = RGBColor(80, 80, 80)
    for label, value in (("Purpose", "Professor discussion brief"), ("Date", "4 September 2026"), ("Status", "Validation evidence available; no final model selected")):
        paragraph = doc.add_paragraph()
        paragraph.paragraph_format.space_after = Pt(2)
        label_run = paragraph.add_run(f"{label}: ")
        label_run.bold = True
        paragraph.add_run(value)

    add_heading(doc, "1. Executive summary")
    doc.add_paragraph(
        "The project aims to identify epochs likely to be light sleep from wearable-style heart-rate and accelerometer data, so a smart alarm can optionally choose a gentler wake moment. The exact user-selected alarm time remains the safety guarantee; the ML component never replaces it."
    )
    add_bullet(doc, "The strongest completed held-out result is the causal-time engineered logistic model: test F1 0.6409 and test AUC 0.6067.")
    add_bullet(doc, "Detailed raw signals alone did not improve the CNN-GRU model. A validation-only raw logistic baseline did retain useful signal.")
    add_bullet(doc, "The new hybrid model gives a small validation improvement (F1 0.6350 versus 0.6331), but the difference is too small to call a confirmed improvement without the professor's guidance on uncertainty and final evaluation.")
    add_bullet(doc, "A high-precision alarm policy selected on validation did not generalize to test participants, showing that conservative thresholding alone is not reliable.")

    add_heading(doc, "2. Data and prediction task")
    doc.add_paragraph("Source: official BIDSleep release 1.0.0. The raw files remain outside the repository; derived artifacts are stored locally under ml/artifacts.")
    add_table(doc, ["Item", "Details"], [
        ("Participants / nights", "47 participants; 253 nights"),
        ("Usable labelled epochs", "213,387 fixed 30-second epochs"),
        ("Split (seed 20260821)", "29 train / 9 validation / 9 held-out test participants"),
        ("Target", "Light sleep (N1/N2) = 1; awake, deep, and REM = 0"),
        ("Raw input", "30 one-second causal bins: acceleration x/y/z + availability; heart rate + availability"),
        ("Coverage finding", "Acceleration is available in about 96% of train bins; heart rate in about 19%"),
    ], [1.7, 4.8])
    doc.add_paragraph("All model fitting and threshold selection use training and validation participants only. The held-out test participants are reserved for a single final evaluation of a candidate selected on validation.")

    add_heading(doc, "3. What has been tried")
    add_table(doc, ["Approach", "Input / design", "Key result", "Interpretation"], [
        ("Causal-time logistic", "12 engineered epoch features plus elapsed-session and clock-time encodings", "Val F1 0.6331; test F1 0.6409; test AUC 0.6067", "Current completed benchmark"),
        ("Conservative policy", "High threshold chosen to target precision >= 0.70 on validation", "Val precision 0.738; test precision 0.486; 107 test triggers", "Did not generalize"),
        ("Rolling engineered", "Causal ten-epoch summaries and rolling movement/HR statistics", "Val F1 0.6314; test F1 0.6399; test AUC 0.6155", "No F1 gain over baseline"),
        ("Raw CNN-GRU", "Current 30-second raw input, 30 x 6, hidden size 64", "Test F1 0.6234; test AUC 0.5146", "Worse than engineered baseline"),
        ("Raw logistic diagnostic", "Standardized current raw 30 x 6 input", "Val F1 0.6175; val AUC 0.6304", "Raw signal has some linear information"),
        ("Hybrid validation candidate", "Current raw 30-second signal + ten-epoch engineered/time context", "Val F1 0.6350; val AUC 0.6371", "Promising but unconfirmed"),
    ], [1.35, 2.25, 1.45, 1.45])

    add_heading(doc, "4. Visual comparison")
    doc.add_picture(str(CHART), width=Inches(6.35))
    caption = doc.add_paragraph("Figure 1. Best F1 selected on validation participants. The hybrid result has not been evaluated on the held-out test split.")
    caption.paragraph_format.space_before = Pt(4)
    caption.paragraph_format.space_after = Pt(8)
    for run in caption.runs:
        run.italic = True
        run.font.size = Pt(9)

    add_heading(doc, "5. Diagnostic findings")
    add_bullet(doc, "Raw-data integrity checks passed: raw labels and participant assignments exactly match the prepared epoch artifact, and sensor timestamps overlap the labelled epochs.")
    add_bullet(doc, "The raw CNN-GRU learned weakly during the diagnostic run: validation AUC reached 0.5477 after 12 epochs; mean loss changed only from 0.7532 to 0.7493; output probabilities were narrowly concentrated.")
    add_bullet(doc, "The hybrid candidate uses 303 standardized features: 180 values from the current raw 30-second epoch, 120 engineered values from ten causal epochs, and three causal time features.")
    add_bullet(doc, "The hybrid's gain over the causal-time baseline is 0.0019 validation F1. It may be real or split-specific variation; it should not be presented as a confirmed improvement yet.")

    add_heading(doc, "6. Limitations and safety")
    add_bullet(doc, "The label is sleep stage, not a direct measurement of how refreshed a person feels after waking. The present study measures epoch classification, not wake-window satisfaction.")
    add_bullet(doc, "BIDSleep participants and wearable sensors may not match the eventual phone/watch user population. External validation is still required.")
    add_bullet(doc, "Heart-rate availability is sparse in the one-second representation, which may limit raw signal models.")
    add_bullet(doc, "The exact-time fallback alarm remains mandatory. The model should be advisory and subject to sensor-quality gates.")

    add_heading(doc, "7. Questions for professor guidance")
    questions = [
        "Is binary light-versus-other sleep stage an appropriate first proxy for a smart alarm, or should the research target be reframed around a wake-window outcome?",
        "How should uncertainty be reported for participant-held-out results: participant-level bootstrap intervals, repeated group splits, or another protocol?",
        "Is the 0.0019 validation-F1 hybrid gain sufficient to justify one final held-out test evaluation, or should more validation-only analysis be required first?",
        "Would a simpler interpretable model be preferable for this project, given the raw CNN-GRU did not add value?",
        "What real-user study, safety criteria, and fairness checks would be needed before product claims are appropriate?",
    ]
    for number, question in enumerate(questions, 1):
        doc.add_paragraph(question, style="List Number")

    add_heading(doc, "8. Proposed next steps")
    add_bullet(doc, "Perform validation-only participant-level uncertainty and calibration analysis for the causal-time baseline and hybrid candidate.")
    add_bullet(doc, "After methodological review, choose at most one candidate and evaluate it once on the held-out participants.")
    add_bullet(doc, "Retain the simple causal-time model if the hybrid advantage is not robust or not worth the added complexity.")
    add_bullet(doc, "Plan a separate real-world evaluation focused on wake-window experience, sensor quality, and alarm safety.")

    add_heading(doc, "9. Reproducibility record")
    doc.add_paragraph("Environment: C:\\Users\\AMIT\\anaconda3\\envs\\fashion-trends\\python.exe. Random seed: 20260821. Key derived artifacts: epochs.parquet, raw_sequences.npz, causal_time_features_validation.json, causal_rolling_features_validation.json, conservative_policy_evaluation.json, raw_cnn_gru/metrics.json, raw_cnn_gru_diagnostic.json, and hybrid_raw_engineered_validation.json.")
    doc.add_paragraph("Code added during the raw/hybrid investigation: ml/raw_sequences.py, ml/diagnose_raw.py, and ml/hybrid_candidate.py, with corresponding focused tests. The ML suite passed 31 tests with three pre-existing ONNX tracer warnings at the time of this report.")

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)


if __name__ == "__main__":
    build_report()
